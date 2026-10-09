"""Exportación del reporte de negocio a Excel (regla 55).

El encabezado lleva el nombre de la barbería, que es dato de `Barbershop` (regla 2).
Todo texto se guarda como celda de texto: un nombre de cliente o una descripción
como "=HYPERLINK(...)" se ve tal cual y nunca se evalúa como fórmula.
"""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from apps.reports.excel.sheets import SHEETS, Sheet
from apps.reports.scope import Period, resolve_period
from apps.tenancy.models import Membership
from apps.tenancy.permissions import ensure_permission

MONEY_FORMAT = "#,##0.00"
DATETIME_FORMAT = "yyyy-mm-dd hh:mm"


def _cell_value(value, period: Period):
    if isinstance(value, datetime):
        # Excel no guarda zona horaria: se escribe la hora local de la barbería.
        return value.astimezone(period.tz).replace(tzinfo=None)
    return value


def _append(sheet: Worksheet, values: list, period: Period) -> None:
    sheet.append([_cell_value(value, period) for value in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            # openpyxl toma "=..." como fórmula; el tipo texto lo impide para cualquier prefijo.
            cell.data_type = "s"
        elif isinstance(cell.value, Decimal):
            cell.number_format = MONEY_FORMAT
        elif isinstance(cell.value, datetime):
            cell.number_format = DATETIME_FORMAT


def _write(book: Workbook, content: Sheet, shop_name: str, period: Period) -> None:
    sheet = book.create_sheet(content.title)
    _append(sheet, [shop_name], period)
    sheet.append([f"{content.title} del {period.date_from:%Y-%m-%d} al {period.date_to:%Y-%m-%d}"])
    sheet.append([])
    sheet.append(content.header)
    for cell in sheet[1] + sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for row in content.rows:
        _append(sheet, row, period)


def build_business_report(membership: Membership, date_from: date | None = None, date_to: date | None = None) -> bytes:
    ensure_permission(membership, "reportes.exportar")
    period = resolve_period(membership, date_from, date_to)
    book = Workbook()
    book.remove(book.active)
    for build_sheet in SHEETS:
        _write(book, build_sheet(membership, period), membership.barbershop.trade_name, period)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
