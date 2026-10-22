"""Página del tablero: permisos, período, exportación, aislamiento y consultas acotadas."""

import re
from datetime import datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.db import connection
from django.test import Client
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from apps.cash import services as cash
from apps.core.tenant_context import tenant_context
from apps.inventory import services as inventory
from apps.sales import services as sales
from apps.sales.selectors import local_today
from apps.staff.models import BarberBranch
from apps.tenancy.models import Branch
from apps.tenancy.roles import Role
from tests.conftest import scope_of
from tests.frontend.conftest import logged_in_with_mfa
from tests.frontend.html_support import elements, forms
from tests.reports.helpers import close_at, sell

pytestmark = pytest.mark.django_db

MONEY_TABLES = ("cash_", "receivables_", "sales_payment")


@pytest.fixture
def owner_client(world):
    return logged_in_with_mfa(world.owner.user)


def test_el_tablero_muestra_los_datos_reales(world, owner_client):
    sell(world, world.ana, service=world.corte)

    page = owner_client.get(reverse("home"))
    html = page.content.decode()

    assert page.context["state"] == "ready"
    assert "Ingresos de hoy" in html and "$25.000" in html
    assert "Sin referencia" in html  # la semana anterior no vendió a esta hora
    assert "sales-calendar" in html and "sales-chart-data" not in html  # la tabla es la vista por defecto


def _view_links(html: str) -> dict:
    return {
        attrs.get("href"): attrs.get("aria-current")
        for tag, attrs, _ in elements(html)
        if tag == "a" and "view-switch__link" in attrs.get("class", "")
    }


def test_la_vista_3d_va_en_la_url_y_trae_los_datos_de_las_torres(world, owner_client):
    sell(world, world.ana, service=world.corte)
    month = f"{local_today(world.owner):%Y-%m}"

    html = owner_client.get(reverse("home"), {"mes": month, "vista": "3d"}).content.decode()

    assert '<script id="sales-chart-data" type="application/json">' in html
    assert "js/modules/sales-towers.js" in html
    assert _view_links(html) == {f"?mes={month}&vista=tabla": None, f"?mes={month}&vista=3d": "page"}
    # Al cambiar de mes se conserva la vista elegida.
    month_form = next(f for f in forms(html) if f["method"] == "get")
    assert "vista" in month_form["fields"] and 'name="vista" value="3d"' in html


@pytest.mark.parametrize("value", ["", "tabla", "torres", "<script>"])
def test_cualquier_otra_vista_es_la_tabla(world, owner_client, value):
    sell(world, world.ana, service=world.corte)

    page = owner_client.get(reverse("home"), {"vista": value})

    assert page.context["view"] == "tabla"
    assert "sales-chart-data" not in page.content.decode()


def test_los_pasos_en_cero_de_la_cascada_van_sin_signo(world, owner_client):
    html = owner_client.get(reverse("home")).content.decode()  # caja abierta, sin movimientos
    amounts = re.findall(r'cascade__amount num">([^<]*)<', html)
    assert amounts and "$0" in amounts
    assert not [amount for amount in amounts if amount in ("+$0", "−$0")]


def test_el_barbero_ve_solo_sus_cortes_y_ninguna_cifra_de_dinero(world):
    sell(world, world.ana, service=world.corte)
    client = Client()
    client.force_login(world.ana_login.user)

    with CaptureQueriesContext(connection) as queries:
        page = client.get(reverse("home"))
    html = page.content.decode()

    assert page.context["state"] == "cuts" and "Mis cortes de hoy" in html
    assert re.search(r'kpi__value num">\s*1\s*<', html)
    assert "$" not in html
    assert not [q["sql"] for q in queries.captured_queries if any(table in q["sql"] for table in MONEY_TABLES)]


def test_exportar_exige_el_permiso_y_responde_un_xlsx(world, owner_client):
    sell(world, world.ana, service=world.corte)
    cashier = Client()
    cashier.force_login(world.cashier.user)

    assert cashier.get(reverse("reports_export")).status_code == 403
    response = owner_client.get(reverse("reports_export"), {"mes": f"{local_today(world.owner):%Y-%m}"})

    today = local_today(world.owner)
    assert response["Content-Type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    assert response["Content-Disposition"] == (
        f'attachment; filename="scrumcut-{world.shop.slug}-centro-{today.replace(day=1):%Y%m%d}-{today:%Y%m%d}.xlsx"'
    )
    assert load_workbook(BytesIO(response.content)).sheetnames[0] == "Panorama"


def _sales_in_export(client, branch) -> list:
    """La sede ya no viaja en la URL: se elige con el cambio de sede y el Excel sigue al alcance."""
    client.post(reverse("switch_branch"), {"sede": str(branch.public_id)})
    response = client.get(reverse("reports_export"))
    rows = load_workbook(BytesIO(response.content))["Ventas"].iter_rows(min_row=5, values_only=True)
    return sorted(row[4] for row in rows)  # la columna "Ítem", después de "Sede"


def test_exportar_una_sede_no_incluye_las_ventas_de_otra(world, owner_client):
    norte = Branch.objects.create(name="Norte")
    BarberBranch.objects.create(barber=world.beto, branch=norte)  # Beto también trabaja en Norte
    cash.open_session(world.owner, scope_of(world.owner, norte), opening_amount=Decimal("0"))
    inventory.register_entry(
        world.owner, world.cera.public_id, scope=scope_of(world.owner, norte), quantity=1, reason="Compra"
    )
    sell(world, world.ana, service=world.corte)  # en Centro
    sales.quick_register(
        world.owner,
        scope_of(world.owner, norte),
        barber=world.beto.public_id,
        method="efectivo",
        product=world.cera.public_id,
    )

    assert _sales_in_export(owner_client, world.branch) == ["Corte"]
    assert _sales_in_export(owner_client, norte) == ["Cera"]
    response = owner_client.get(reverse("reports_export"))  # sigue en Norte, la última sede elegida
    assert "-norte-" in response["Content-Disposition"]


@pytest.mark.parametrize("value", ["siempre", "2026-13", "2999-01", "1999-01", "2026-10-01'--"])
def test_un_mes_invalido_futuro_o_anterior_a_la_barberia_vuelve_al_mes_actual(world, owner_client, value):
    page = owner_client.get(reverse("home"), {"mes": value})

    assert page.status_code == 200
    assert page.context["dashboard"]["date_from"] == local_today(world.owner).replace(day=1)


def test_un_mes_sin_ventas_sale_en_ceros_y_la_pagina_carga(world, owner_client):
    page = owner_client.get(reverse("home"))
    calendar = page.context["dashboard"]["calendar"]

    assert page.context["state"] == "ready" and not page.context["dashboard"]["has_sales"]
    assert (calendar["total"], calendar["selling_days"]) == (0, 0)
    assert "sales-calendar" in page.content.decode() and "sales-chart-data" not in page.content.decode()


def test_exportar_solo_aparece_con_el_permiso(world, owner_client):
    cashier = Client()
    cashier.force_login(world.cashier.user)

    assert "Exportar a Excel" in owner_client.get(reverse("home")).content.decode()
    assert "Exportar a Excel" not in cashier.get(reverse("home")).content.decode()


def test_cada_barberia_ve_solo_sus_ventas(world, member, other_shop):
    sell(world, world.ana, service=world.corte)
    with tenant_context(other_shop.pk):
        theirs = logged_in_with_mfa(member(Role.OWNER, other_shop).user)

    dashboard = theirs.get(reverse("home")).context["dashboard"]

    assert not dashboard["has_sales"] and dashboard["kpis"]["income"]["value"] == 0


# El marco (sesión, sedes, alertas y resumen de Administración) y el tablero completo de la página de
# inicio del dueño: cada bloque es una consulta agregada, así que el total no depende de sedes ni de días.
HOME_QUERIES = 26


def _branches_with_sales(world, extra: int) -> None:
    for number in range(extra):
        branch = Branch.objects.create(name=f"Sede {number}")
        BarberBranch.objects.create(barber=world.beto, branch=branch)
        inventory.register_entry(
            world.owner, world.cuchilla.public_id, scope=scope_of(world.owner, branch), quantity=2, reason="Compra"
        )
        cash.open_session(world.owner, scope_of(world.owner, branch), opening_amount=Decimal("0"))
        sales.quick_register(
            world.owner,
            scope_of(world.owner, branch),
            barber=world.beto.public_id,
            method="efectivo",
            service=world.corte.public_id,
        )


@pytest.mark.parametrize(("extra_branches", "selling_days"), [(0, 1), (2, 5)])
def test_las_consultas_del_marco_y_del_tablero_no_crecen(
    world, owner_client, django_assert_num_queries, extra_branches, selling_days
):
    """Con 1 sede y 1 día con ventas, o con 3 sedes y 5 días: el mismo número de consultas."""
    today = local_today(world.owner)
    for day in range(selling_days):
        close_at(sell(world, world.ana, service=world.corte), datetime.combine(today - timedelta(days=day), time(9)))
    _branches_with_sales(world, extra_branches)
    owner_client.post(reverse("switch_branch"), {"sede": "todas"})

    with django_assert_num_queries(HOME_QUERIES):
        owner_client.get(reverse("home"))
