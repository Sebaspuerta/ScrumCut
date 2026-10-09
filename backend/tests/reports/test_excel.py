"""Exportación a Excel (regla 55): hojas, marca de la barbería, días sin ventas y texto que nunca es fórmula."""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.cash import services as cash
from apps.reports.excel.builder import build_business_report
from tests.reports.helpers import close_at, sell

pytestmark = pytest.mark.django_db

SHEETS = [
    "Panorama",
    "Ventas",
    "Comisiones",
    "Fiados",
    "Movimientos de caja",
    "Cierres de caja",
    "Estado de resultados",
]
FORMULA_LIKE = ["=1+1", "+57 300", "-3", "@SUM(A1)"]


def _workbook(world, *dates):
    sell(world, world.ana, service=world.corte)
    sell(world, world.beto, product=world.cera)
    for text in FORMULA_LIKE:
        cash.register_manual_movement(
            world.cashier, world.session.public_id, "gasto", Decimal("1000"), "efectivo", text
        )
    return load_workbook(BytesIO(build_business_report(world.owner, *dates)))


def test_el_excel_tiene_las_hojas_y_el_nombre_de_la_barberia(world):
    book = _workbook(world)
    assert book.sheetnames == SHEETS
    assert {book[name]["A1"].value for name in SHEETS} == {"Barbería a"}
    sales = list(book["Ventas"].iter_rows(min_row=5, values_only=True))
    cera = next(row for row in sales if row[3] == "Cera")
    assert (cera[9], cera[10], cera[12]) == (20000, 8000, 12000)  # venta neta, costo y margen


def test_el_texto_que_parece_formula_se_guarda_como_texto(world):
    movements = _workbook(world)["Movimientos de caja"]
    cells = {cell.value: cell for row in movements.iter_rows(min_row=5) for cell in row if cell.value in FORMULA_LIKE}
    assert sorted(cells) == sorted(FORMULA_LIKE)  # sin apóstrofo ni cambios
    assert {cell.data_type for cell in cells.values()} == {"s"}  # texto, nunca "f" (fórmula)


def test_el_panorama_muestra_los_dias_sin_ventas_en_cero(world):
    first, second = sell(world, world.ana, service=world.corte), sell(world, world.beto, service=world.barba)
    close_at(first, datetime(2026, 10, 5, 12, 0))
    close_at(second, datetime(2026, 10, 7, 12, 0))

    overview = load_workbook(BytesIO(build_business_report(world.owner, date(2026, 10, 5), date(2026, 10, 7))))[
        "Panorama"
    ]

    days = [row[:3] for row in overview.iter_rows(values_only=True) if isinstance(row[0], datetime)]
    assert days == [(datetime(2026, 10, 5), 1, 25000), (datetime(2026, 10, 6), 0, 0), (datetime(2026, 10, 7), 1, 15000)]
